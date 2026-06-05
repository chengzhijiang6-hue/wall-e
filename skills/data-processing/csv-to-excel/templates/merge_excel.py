"""
Excel 合并脚本模板 — 将多个已生成的 Excel 合并为一个
用法: 修改 file_a, file_b, output_path 后直接运行
"""
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ===== 配置 =====
file_a = "/path/to/file_a.xlsx"
file_b = "/path/to/file_b.xlsx"
output_path = "/path/to/merged.xlsx"

# ===== 样式 =====
thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
hdr_font = Font(bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor="4472C4")
sum_fill = PatternFill("solid", fgColor="2F5496")
alt_fill = PatternFill("solid", fgColor="D9E2F3")
green = PatternFill("solid", fgColor="C6EFCE")
red = PatternFill("solid", fgColor="FFC7CE")

def fmt_hdr(ws, ncols, fill=hdr_fill):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = hdr_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin

def fmt_data(ws, nrows, ncols, start=2):
    for r in range(start, start + nrows):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = thin
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if (r - start) % 2 == 1:
                cell.fill = alt_fill

def auto_w(ws, ncols, mx=35):
    for c in range(1, ncols + 1):
        ml = max((len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(ml + 4, mx)

def copy_sheet_data(src_ws, dst_ws, start_row=1):
    """复制 sheet 数据（含格式）到目标 sheet"""
    for row in src_ws.iter_rows(min_row=start_row, values_only=False):
        for cell in row:
            dst_ws.cell(cell.row, cell.column, cell.value)
            if cell.row == 1:
                dst_ws.cell(cell.row, cell.column).font = hdr_font
                dst_ws.cell(cell.row, cell.column).fill = hdr_fill
            dst_ws.cell(cell.row, cell.column).border = thin
            dst_ws.cell(cell.row, cell.column).alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True
            )

# ===== 读取 =====
wb_a = load_workbook(file_a)
wb_b = load_workbook(file_b)

# ===== 合并连接明细 =====
# 策略: 增加"方向"列区分来源（如出站/入站）
wb = Workbook()
ws1 = wb.active
ws1.title = "连接明细"

# 读取表头，插入"方向"列
src_headers = [cell.value for cell in wb_a["连接明细"][1]]
merged_headers = [src_headers[0], "方向"] + src_headers[1:]
for c, h in enumerate(merged_headers, 1):
    ws1.cell(1, c, h)

row_idx = 2
# 复制 file_a 数据
for r in range(2, wb_a["连接明细"].max_row + 1):
    ws1.cell(row_idx, 1, row_idx - 1)
    ws1.cell(row_idx, 2, "出站")
    for c in range(1, len(src_headers) + 1):
        val = wb_a["连接明细"].cell(r, c).value
        ws1.cell(row_idx, c + 2, val)
    row_idx += 1

# 复制 file_b 数据
for r in range(2, wb_b["连接明细"].max_row + 1):
    ws1.cell(row_idx, 1, row_idx - 1)
    ws1.cell(row_idx, 2, "入站")
    for c in range(1, len(src_headers) + 1):
        val = wb_b["连接明细"].cell(r, c).value
        ws1.cell(row_idx, c + 2, val)
    row_idx += 1

fmt_hdr(ws1, len(merged_headers))
fmt_data(ws1, row_idx - 2, len(merged_headers))
auto_w(ws1, len(merged_headers))

# ===== 合并汇总表 =====
# 策略: 用 defaultdict 累加同名维度的数值字段
def merge_summary_sheets(ws_a, ws_b, wb_out, sheet_name, direction_a="来源A", direction_b="来源B"):
    """合并两个结构相同的汇总 sheet"""
    headers = [cell.value for cell in ws_a[1]]
    
    # 假设第一列是维度 key，最后一列是命中次数
    key_col = 1
    hits_col = len(headers)
    count_col = 4  # 连接数列（根据实际调整）
    
    merged = defaultdict(lambda: {"data": {}, "hits": 0, "count": 0})
    
    for ws, direction in [(ws_a, direction_a), (ws_b, direction_b)]:
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, key_col).value
            if not key:
                continue
            g = merged[key]
            for c in range(1, len(headers) + 1):
                val = ws.cell(r, c).value
                if c == hits_col:
                    g["hits"] += (val or 0)
                elif c == count_col:
                    g["count"] += (val or 0)
                else:
                    # 合并文本（去重）
                    if val:
                        if c not in g["data"]:
                            g["data"][c] = set()
                        g["data"][c].add(str(val))
    
    ws_out = wb_out.create_sheet(sheet_name)
    for c, h in enumerate(headers, 1):
        ws_out.cell(1, c, h)
    
    row_idx = 2
    for key, g in sorted(merged.items(), key=lambda x: x[1]["hits"], reverse=True):
        ws_out.cell(row_idx, key_col, key)
        for c in range(1, len(headers) + 1):
            if c == hits_col:
                ws_out.cell(row_idx, c, g["hits"])
            elif c == count_col:
                ws_out.cell(row_idx, c, g["count"])
            elif c in g["data"]:
                ws_out.cell(row_idx, c, "\n".join(sorted(g["data"][c])))
        ws_out.row_dimensions[row_idx].height = 35
        row_idx += 1
    
    fmt_hdr(ws_out, len(headers), sum_fill)
    fmt_data(ws_out, row_idx - 2, len(headers))
    auto_w(ws_out, len(headers))

# 使用示例：
# merge_summary_sheets(wb_a["按防火墙汇总"], wb_b["按防火墙汇总"], wb, "按防火墙汇总")

# ===== 保存 =====
wb.save(output_path)
print(f"合并完成: {output_path}")

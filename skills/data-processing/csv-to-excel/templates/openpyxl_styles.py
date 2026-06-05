"""
openpyxl 样式函数库 — 从 csv-to-excel 技能中提取的通用格式化函数
用法: 直接复制到脚本顶部或 import
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# 通用样式
thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
hdr_font = Font(bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor="4472C4")
sum_fill = PatternFill("solid", fgColor="2F5496")
alt_fill = PatternFill("solid", fgColor="D9E2F3")
green  = PatternFill("solid", fgColor="C6EFCE")
red    = PatternFill("solid", fgColor="FFC7CE")
orange = PatternFill("solid", fgColor="FFEB9C")

def fmt_hdr(ws, ncols, fill=hdr_fill):
    """格式化表头行"""
    for c in range(1, ncols+1):
        cell = ws.cell(1, c)
        cell.font = hdr_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin

def fmt_data(ws, nrows, ncols, start=2):
    """格式化数据区：隔行着色 + 边框"""
    for r in range(start, start+nrows):
        for c in range(1, ncols+1):
            cell = ws.cell(r, c)
            cell.border = thin
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if (r - start) % 2 == 1:
                cell.fill = alt_fill

def auto_w(ws, ncols, mx=35):
    """自动列宽"""
    for c in range(1, ncols+1):
        ml = max((len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row+1)), default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(ml + 4, mx)

def copy_sheet_data(src_ws, dst_ws, start_row=1):
    """复制 sheet 数据到目标 sheet（含表头样式）"""
    for row in src_ws.iter_rows(min_row=start_row, values_only=False):
        for cell in row:
            dst_ws.cell(cell.row, cell.column, cell.value)
            if cell.row == 1:
                dst_ws.cell(cell.row, cell.column).font = hdr_font
                dst_ws.cell(cell.row, cell.column).fill = hdr_fill
            dst_ws.cell(cell.row, cell.column).border = thin
            dst_ws.cell(cell.row, cell.column).alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True
            )

def style_risk_cell(cell, risk):
    """风险等级着色：HIGH红/MEDIUM橙/LOW绿"""
    if risk == 'HIGH': cell.fill = red
    elif risk == 'MEDIUM': cell.fill = orange
    else: cell.fill = green

def aggregate_by_key(data, key_col, headers):
    """按指定列分组聚合（用于多维度汇总）
    
    返回: dict[key] -> {count, protocols, ports, services, actions, total_hits}
    """
    groups = defaultdict(lambda: {
        'count': 0,
        'protocols': set(),
        'ports': set(),
        'services': set(),
        'actions': set(),
        'total_hits': 0
    })
    
    # 自动检测列索引
    proto_idx = headers.index('Protocol') if 'Protocol' in headers else -1
    port_idx = headers.index('Port') if 'Port' in headers else -1
    svc_idx = headers.index('Service_Name') if 'Service_Name' in headers else -1
    action_idx = headers.index('Action') if 'Action' in headers else -1
    hits_idx = headers.index('Hits') if 'Hits' in headers else -1
    
    for row in data:
        key = row[key_col] if key_col < len(row) else ''
        g = groups[key]
        g['count'] += 1
        
        if proto_idx >= 0 and proto_idx < len(row) and row[proto_idx]: 
            g['protocols'].add(row[proto_idx])
        if port_idx >= 0 and port_idx < len(row) and row[port_idx]: 
            g['ports'].add(row[port_idx])
        if svc_idx >= 0 and svc_idx < len(row) and row[svc_idx]: 
            g['services'].add(row[svc_idx])
        if action_idx >= 0 and action_idx < len(row) and row[action_idx]: 
            g['actions'].add(row[action_idx])
        if hits_idx >= 0 and hits_idx < len(row):
            try: g['total_hits'] += int(row[hits_idx])
            except: pass
    
    return groups

def create_summary_sheet(wb, name, groups, key_name, extra_headers=None):
    """创建汇总Sheet
    
    Args:
        wb: Workbook对象
        name: Sheet名称
        groups: aggregate_by_key返回的字典
        key_name: 主键列名（如"目标IP"）
        extra_headers: 额外列名列表（默认为['协议','端口','服务/应用','动作']）
    """
    if extra_headers is None:
        extra_headers = ['协议', '端口', '服务/应用', '动作']
    
    ws = wb.create_sheet(name)
    headers = [key_name, '连接数'] + extra_headers + ['总命中次数']
    ws.append(headers)
    
    # 按总命中次数降序排列
    sorted_groups = sorted(groups.items(), key=lambda x: x[1]['total_hits'], reverse=True)
    
    for key, g in sorted_groups:
        row = [
            key,
            g['count'],
            '\n'.join(sorted(g['protocols'])) or '-',
            '\n'.join(sorted(g['ports'])) or '-',
            '\n'.join(sorted(g['services'])) or '-',
            '\n'.join(sorted(g['actions'])) or '-',
            g['total_hits']
        ]
        ws.append(row)
    
    # 格式化
    fmt_hdr(ws, len(headers))
    fmt_data(ws, len(sorted_groups), len(headers))
    auto_w(ws, len(headers))
    
    # 多值字段设置自动换行和行高
    for r in range(2, ws.max_row+1):
        max_lines = 1
        for c in range(3, len(headers)):  # 协议、端口、服务、动作列
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='center')
            lines = str(ws.cell(r, c).value or '').count('\n') + 1
            max_lines = max(max_lines, lines)
        ws.row_dimensions[r].height = max(30, 15 * max_lines)
    
    return ws

# 多 CSV 合并为多 Sheet Excel（带非微软高亮）

```python
import csv, os, sys, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Config: update device IP as needed
DEVICE_IP = "10.177.104.122"
BASE_DIR = r"C:\Users\CZE8WX\Desktop\work\server-migration"
CSV_DIR = os.path.join(BASE_DIR, DEVICE_IP, "raw")
OUTPUT = os.path.join(BASE_DIR, DEVICE_IP, "merged", "Migration_Report.xlsx")

CSV_SHEETS = [
    ("1_Port_App_Mapping.csv", "Port_App_Mapping"),
    ("2_All_Services.csv", "All_Services"),
    ("3_Scheduled_Tasks.csv", "Scheduled_Tasks"),
    ("4_Installed_Software.csv", "Installed_Software"),
    ("5_Network_Shares.csv", "Network_Shares"),
    ("6_IIS_Sites.csv", "IIS_Sites"),
    ("7_ODBC_DataSources.csv", "ODBC_DataSources"),
    ("8a_Env_Variables.csv", "Env_Variables"),
]

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
DATA_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
ALT_FILL = PatternFill("solid", fgColor="D9E2F3")
HIGHLIGHT = PatternFill("solid", fgColor="FFC7CE")

def auto_width(ws, mx=40):
    for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        ml = max((len(str(c.value or "")) for c in col_cells), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(ml + 4, mx)

wb = Workbook(); wb.remove(wb.active)
for csv_name, sheet_name in CSV_SHEETS:
    ws = wb.create_sheet(title=sheet_name)
    csv_path = os.path.join(CSV_DIR, csv_name)
    if not os.path.exists(csv_path) or os.path.getsize(csv_path) <= 3:
        ws.cell(1, 1, "No Data"); ws.cell(1, 1).font = HEADER_FONT; ws.cell(1, 1).fill = HEADER_FILL
        continue
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if not rows: continue
    headers = rows[0]
    ms_idx = headers.index("Is_Microsoft") + 1 if "Is_Microsoft" in headers else None
    default_idx = headers.index("Is_Default") + 1 if "Is_Default" in headers else None
    for c, val in enumerate(headers, 1):
        cell = ws.cell(1, c, value=val); cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN; cell.border = THIN_BORDER
    for r, row in enumerate(rows[1:], 2):
        highlight = False
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, value=val); cell.alignment = DATA_ALIGN; cell.border = THIN_BORDER
            if ms_idx and c == ms_idx and val == "No": highlight = True
            if default_idx and c == default_idx and val == "No": highlight = True
        if highlight:
            for c in range(1, len(row)+1): ws.cell(r, c).fill = HIGHLIGHT
        elif (r-2) % 2 == 1:
            for c in range(1, len(row)+1): ws.cell(r, c).fill = ALT_FILL
    ws.freeze_panes = "A2"; auto_width(ws)
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
wb.save(OUTPUT)
print(f"Done: {OUTPUT}")
```

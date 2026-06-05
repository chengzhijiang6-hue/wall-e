# 非微软项目分析脚本（依赖关系 + 风险评估 + 迁移清单）

```python
# -*- coding: utf-8 -*-
"""
Non-Microsoft Items Analysis Script
Reads CSV files from raw/, filters non-MS items, analyzes dependencies
Outputs: Non_MS_Items_Report.xlsx to merged/
"""
import os, sys, csv, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference
from collections import defaultdict

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

# Config: update device IP as needed
DEVICE_IP = "10.177.104.122"
BASE_DIR = r"C:\Users\CZE8WX\Desktop\work\server-migration"
CSV_DIR = os.path.join(BASE_DIR, DEVICE_IP, "raw")
OUTPUT_FILE = os.path.join(BASE_DIR, DEVICE_IP, "merged", "Non_MS_Items_Report.xlsx")

# Styles
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
DATA_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
ALT_FILL = PatternFill("solid", fgColor="D9E2F3")
HIGH_RISK_FILL = PatternFill("solid", fgColor="FFC7CE")
MEDIUM_RISK_FILL = PatternFill("solid", fgColor="FFEB9C")
LOW_RISK_FILL = PatternFill("solid", fgColor="C6EFCE")

HIGH_RISK_PORTS = {21, 22, 23, 25, 53, 80, 110, 135, 139, 143, 443, 445, 993, 995, 1433, 1521, 3306, 3389, 5432, 5900, 5985, 8080, 8443}
MEDIUM_RISK_PORTS = {20, 69, 88, 161, 162, 389, 636, 1080, 1723, 2049, 3268, 3269, 5060, 5061, 5986, 7001, 8000, 8001, 8008, 8888, 9090}


def read_csv(filepath):
    if not os.path.exists(filepath): return []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def auto_width(ws, max_width=40):
    for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        ml = max((len(str(c.value or "")) for c in col_cells), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(ml + 4, max_width)


def write_header(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN; cell.border = THIN_BORDER


def write_row(ws, row_num, data, risk_level=None):
    for c, val in enumerate(data, 1):
        cell = ws.cell(row_num, c, value=val)
        cell.alignment = DATA_ALIGN; cell.border = THIN_BORDER
        if risk_level == "high": cell.fill = HIGH_RISK_FILL
        elif risk_level == "medium": cell.fill = MEDIUM_RISK_FILL
        elif (row_num - 2) % 2 == 1: cell.fill = ALT_FILL


def get_port_risk(port):
    try:
        p = int(port)
        if p in HIGH_RISK_PORTS: return "high"
        elif p in MEDIUM_RISK_PORTS: return "medium"
    except: pass
    return "low"


def main():
    print("=" * 60)
    print("  Non-Microsoft Items Analysis")
    print("=" * 60)

    ports = read_csv(os.path.join(CSV_DIR, "1_Port_App_Mapping.csv"))
    services = read_csv(os.path.join(CSV_DIR, "2_All_Services.csv"))
    tasks = read_csv(os.path.join(CSV_DIR, "3_Scheduled_Tasks.csv"))
    software = read_csv(os.path.join(CSV_DIR, "4_Installed_Software.csv"))
    shares = read_csv(os.path.join(CSV_DIR, "5_Network_Shares.csv"))

    non_ms_ports = [p for p in ports if p.get("Is_Microsoft", "Yes") == "No"]
    non_ms_services = [s for s in services if s.get("Is_Microsoft", "Yes") == "No"]
    non_ms_tasks = [t for t in tasks if t.get("Is_Microsoft", "Yes") == "No"]
    non_ms_software = [sw for sw in software if sw.get("Is_Microsoft", "Yes") == "No"]
    non_ms_shares = [sh for sh in shares if sh.get("Is_Default", "Yes") == "No"]

    # Build dependency maps
    svc_port_map = defaultdict(set)
    for p in ports:
        sn = p.get("Service_Name", "")
        if sn and sn != "None": svc_port_map[sn].add(p.get("Port", ""))

    wb = Workbook(); wb.remove(wb.active)

    # Dashboard
    ws = wb.create_sheet(title="Dashboard")
    data = [
        ["Category", "Total", "Non-Microsoft", "Percentage"],
        ["Ports", len(ports), len(non_ms_ports), f"{len(non_ms_ports)/max(len(ports),1)*100:.1f}%"],
        ["Services", len(services), len(non_ms_services), f"{len(non_ms_services)/max(len(services),1)*100:.1f}%"],
        ["Tasks", len(tasks), len(non_ms_tasks), f"{len(non_ms_tasks)/max(len(tasks),1)*100:.1f}%"],
        ["Software", len(software), len(non_ms_software), f"{len(non_ms_software)/max(len(software),1)*100:.1f}%"],
        ["Shares", len(shares), len(non_ms_shares), f"{len(non_ms_shares)/max(len(shares),1)*100:.1f}%"],
        ["", "", "", ""],
        ["TOTAL", len(non_ms_ports)+len(non_ms_services)+len(non_ms_tasks)+len(non_ms_software)+len(non_ms_shares), "", ""],
    ]
    for r, row in enumerate(data, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, value=val)
            if r == 1: cell.font = HEADER_FONT; cell.fill = HEADER_FILL
            cell.alignment = DATA_ALIGN; cell.border = THIN_BORDER
    auto_width(ws)

    # Ports
    ws = wb.create_sheet(title="Ports")
    write_header(ws, ["Port","Protocol","Identity","Service","Path","PID","Publisher","System_Port","Risk"])
    for r, p in enumerate(non_ms_ports, 2):
        risk = get_port_risk(p.get("Port",""))
        write_row(ws, r, [p.get("Port",""),p.get("Protocol",""),p.get("Identity",""),p.get("Service_Name",""),p.get("Installation_Path",""),p.get("Process_ID",""),p.get("Publisher",""),p.get("Is_System_Port",""),risk.upper()], risk)
    auto_width(ws)

    # Services
    ws = wb.create_sheet(title="Services")
    write_header(ws, ["Name","DisplayName","State","StartMode","PID","Account","Path","Publisher","Ports"])
    for r, s in enumerate(non_ms_services, 2):
        sn = s.get("Name","")
        rp = ", ".join(svc_port_map.get(sn, set()))
        write_row(ws, r, [sn,s.get("DisplayName",""),s.get("State",""),s.get("StartMode",""),s.get("ProcessId",""),s.get("LogOnAccount",""),s.get("PathName",""),s.get("Publisher",""),rp])
    auto_width(ws)

    # Tasks
    ws = wb.create_sheet(title="Tasks")
    write_header(ws, ["TaskName","TaskPath","State","Author","LastRun","NextRun","Action"])
    for r, t in enumerate(non_ms_tasks, 2):
        write_row(ws, r, [t.get("TaskName",""),t.get("TaskPath",""),t.get("State",""),t.get("Author",""),t.get("LastRunTime",""),t.get("NextRunTime",""),t.get("Action","")])
    auto_width(ws)

    # Software
    ws = wb.create_sheet(title="Software")
    write_header(ws, ["Name","Version","Publisher","InstallDate","Path","Services","Ports"])
    for r, sw in enumerate(non_ms_software, 2):
        name = sw.get("DisplayName","")
        svcs = [s.get("Name","") for s in non_ms_services if name.lower() in s.get("PathName","").lower()]
        ports_set = set()
        for svc in svcs: ports_set.update(svc_port_map.get(svc, set()))
        write_row(ws, r, [name,sw.get("DisplayVersion",""),sw.get("Publisher",""),sw.get("InstallDate",""),sw.get("InstallLocation",""),", ".join(svcs),", ".join(ports_set)])
    auto_width(ws)

    # Shares
    ws = wb.create_sheet(title="Shares")
    write_header(ws, ["Name","Path","Description"])
    for r, sh in enumerate(non_ms_shares, 2):
        write_row(ws, r, [sh.get("Name",""),sh.get("Path",""),sh.get("Description","")])
    auto_width(ws)

    # Dependencies
    ws = wb.create_sheet(title="Dependencies")
    write_header(ws, ["Source_Type","Source_Name","Target_Type","Target_Name","Relationship"])
    dr = 2
    for svc in non_ms_services:
        sn = svc.get("Name","")
        for port in svc_port_map.get(sn, set()):
            write_row(ws, dr, ["Service",sn,"Port",port,"Listens_On"]); dr += 1
        pid = svc.get("ProcessId","")
        if pid: write_row(ws, dr, ["Service",sn,"Process",f"PID:{pid}","Runs_As"]); dr += 1
    for task in non_ms_tasks:
        action = task.get("Action","")
        for sw in non_ms_software:
            if sw.get("DisplayName","").lower() in action.lower():
                write_row(ws, dr, ["Task",task.get("TaskName",""),"Software",sw.get("DisplayName",""),"Executes"]); dr += 1
    auto_width(ws)

    # Migration Checklist
    ws = wb.create_sheet(title="Migration_Checklist")
    write_header(ws, ["Item_Name","Type","Risk","Ports","Dependencies","Action","Status"])
    cr = 2
    for p in non_ms_ports:
        risk = get_port_risk(p.get("Port",""))
        write_row(ws, cr, [f"Port {p.get('Port','')}","Port",risk.upper(),p.get("Port",""),p.get("Service_Name",""),"Verify port needed",""], risk); cr += 1
    for s in non_ms_services:
        rp = ", ".join(svc_port_map.get(s.get("Name",""), set()))
        write_row(ws, cr, [s.get("Name",""),"Service","MEDIUM",rp,"",f"Check: {s.get('State','')}",""]); cr += 1
    for t in non_ms_tasks:
        write_row(ws, cr, [t.get("TaskName",""),"Task","LOW","","",f"Path: {t.get('TaskPath','')}",""]); cr += 1
    for sw in non_ms_software:
        write_row(ws, cr, [sw.get("DisplayName",""),"Software","LOW","","","Check license",""]); cr += 1
    for sh in non_ms_shares:
        write_row(ws, cr, [sh.get("Name",""),"Share","LOW","","",f"Path: {sh.get('Path','')}",""]); cr += 1
    auto_width(ws)

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(f"\nDone: {OUTPUT_FILE}")
    print(f"  Non-MS: {len(non_ms_ports)} ports, {len(non_ms_services)} services, {len(non_ms_tasks)} tasks, {len(non_ms_software)} software, {len(non_ms_shares)} shares")
    print(f"  Dependencies: {dr-2} relationships")
    print(f"  Checklist: {cr-2} items")

if __name__ == "__main__":
    main()
```
